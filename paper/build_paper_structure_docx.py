from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "paper/ERL_Paper_Overall_Structure_Cenergia.docx"

NAVY = "17365D"
BLUE = "2F5597"
LIGHT_BLUE = "D9EAF7"
PALE = "EEF3F8"
DARK = "263238"
GRAY = "666666"
WHITE = "FFFFFF"
GREEN = "E2F0D9"
ORANGE = "FCE4D6"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=80, start=100, bottom=80, end=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def cant_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def font(run, size=10.5, bold=None, color=DARK, italic=None, name="Lato"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def para(text="", size=10.5, bold=False, color=DARK, italic=False,
         before=0, after=4, line=1.08, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    font(r, size=size, bold=bold, color=color, italic=italic)
    return p


def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    font(r, size=14, bold=True, color=NAVY)
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    font(p.add_run(text), size=10.2)
    return p


def number(text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    font(p.add_run(text), size=10.2)
    return p


def callout(title, body, fill=LIGHT_BLUE):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    shade(cell, fill)
    margins(cell, top=120, start=160, bottom=120, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    font(p.add_run(title), size=10.5, bold=True, color=NAVY)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.05
    font(p2.add_run(body), size=10.2)
    return t


for style_name in ["Normal", "List Bullet", "List Number"]:
    st = doc.styles[style_name]
    st.font.name = "Lato"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Lato")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Lato")
    st.font.size = Pt(10.5)

# Header and footer
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(hp.add_run("CENERGIA LAB  |  ERL MANUSCRIPT STRUCTURE"), size=8, bold=True, color=GRAY)

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(fp.add_run("From NDC pathways to production networks  |  Working paper structure  |  September 2026"),
     size=8, color=GRAY)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(26)
p.paragraph_format.space_after = Pt(10)
font(p.add_run("Overall Structure of the ERL Paper"), size=24, bold=True, color=NAVY)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
font(p.add_run("From NDC pathways to production networks: Macroeconomic spillovers of Brazil’s net-zero transition"),
     size=15, bold=True, color=BLUE)

para("Target journal: Environmental Research Letters (ERL)", bold=True, after=2)
para("Project: Cenergia / Brazil climate-macro production-network model", after=2)
para("Purpose of this document: concise manuscript architecture for project reporting and paper development.",
     size=10.2, color=GRAY, italic=True, after=12)

callout(
    "Central research question",
    "How do NDC-consistent investment and technology pathways propagate through Brazil’s production network, "
    "and how does enforcing consistency with sectoral physical production pathways alter the resulting "
    "macroeconomic and distributional effects?"
)
doc.add_paragraph().paragraph_format.space_after = Pt(4)
callout(
    "Core framing",
    "NDC implementation is not only an emissions pathway or an investment shock. The paper treats it as a "
    "restructuring of production networks driven by transition investment and changes in intermediate energy use, "
    "while selected sectoral outputs remain anchored to the physical NDC pathway.",
    fill=PALE
)

heading("1. Paper objective and contribution")
para(
    "The paper soft-links Brazilian climate-planning pathways (MMA/BLUES) with the 73-sector energy-disaggregated "
    "2018 input-output framework developed by EPE/FIPE. The aim is to translate technology, investment and physical "
    "production pathways into production-network consequences that can be expressed in gross output, value added, "
    "sectoral activity and labor requirements.",
    after=5
)
for item in [
    "Connect detailed NDC implementation pathways to inter-industry production networks rather than using stylized green-investment shocks alone.",
    "Separate two economic transmission channels from a physical consistency layer, avoiding the interpretation of the physical pathway as a third independent demand shock.",
    "Show how domestic-content assumptions and supply-chain linkages determine how much transition expenditure is retained in Brazil.",
    "Provide a transparent, reproducible framework that can be transferred to other emerging-economy NDC assessments."
]:
    bullet(item)

heading("2. Conceptual architecture of the model")
t = doc.add_table(rows=2, cols=3)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
heads = [
    "Channel 1\nTransition investment",
    "Channel 2\nIntermediate energy substitution",
    "Physical consistency layer\nNDC production constraints"
]
bodies = [
    "Transition expenditure is reconstructed over time and allocated to domestic capital-goods, construction, "
    "infrastructure, engineering and related supplier sectors. Imports are removed through domestic-content factors.",
    "Scenario-specific physical energy shares rewire selected technical coefficients in A, redistributing intermediate "
    "demand among fossil fuels, bioenergy and electricity at constant base-year valuation.",
    "Selected sectoral outputs are constrained to match external MMA/BLUES physical trajectories. The resulting "
    "balancing final-demand term is a consistency residual, not an autonomous economic shock."
]
for j, h in enumerate(heads):
    c = t.cell(0, j)
    shade(c, NAVY)
    margins(c)
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(h), size=9.4, bold=True, color=WHITE)
for j, b in enumerate(bodies):
    c = t.cell(1, j)
    margins(c, top=120, bottom=120)
    p = c.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    font(p.add_run(b), size=9.1)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(5)
p.paragraph_format.line_spacing = 1.06
font(p.add_run("Baseline/counterfactual: "), size=10.5, bold=True, color=NAVY)
font(p.add_run(
    "All sectors follow the exogenous GDP path with the 2018 technical-coefficient matrix fixed. No transition-specific "
    "investment, energy-input rewiring or NDC production constraint is imposed. Reported differences are scenario "
    "differentials, not causal policy estimates."
), size=10.3)

heading("3. Main manuscript structure")
t = doc.add_table(rows=1, cols=3)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
repeat_header(t.rows[0])
for j, h in enumerate(["Section", "Subsection", "Purpose / content"]):
    c = t.cell(0, j)
    shade(c, NAVY)
    margins(c)
    font(c.paragraphs[0].add_run(h), size=9.4, bold=True, color=WHITE)

rows = [
    ("1. Introduction", "Motivation and gap", "Why NDC implementation creates economy-wide effects through investment, technology choices, intermediate inputs and supply chains; why Brazil is a distinctive case."),
    ("1. Introduction", "Contribution and question", "Position the paper as a soft link between climate/energy pathways and a national production network; state the central research question and the paper’s main contribution."),
    ("2. Methods", "2.1 Data and sectoral representation", "Describe MMA/BLUES pathways and the 73-sector MIP-EPE 2018 system, including the open Type-I Leontief formulation and satellite accounts."),
    ("2. Methods", "2.2 Scenario design and baseline", "Define the all-sector GDP-growth, fixed-structure counterfactual and the transition differential."),
    ("2. Methods", "2.3 Channel 1: transition investment", "Reconstruct investment waves; map technologies to purchaser/supplier sectors; apply domestic-content/import-leakage factors."),
    ("2. Methods", "2.4 Channel 2: energy-input substitution", "Modify selected technical coefficients using MMA/BLUES physical energy shares. For normalized blocks, hold the aggregate monetary energy-input coefficient at its base IO value while changing composition."),
    ("2. Methods", "2.5 Physical consistency layer", "Impose sector-specific NDC production constraints and calculate the residual balancing final demand needed to reconcile the IO solution with the external physical pathway."),
    ("2. Methods", "2.6 Outcomes and satellites", "Calculate gross output, value added and labor requirements using 2018 satellite coefficients; explain that employment is a job-equivalent/labor-requirement measure, not a forecast."),
    ("2. Methods", "2.7 Decomposition and sensitivities", "Distinguish economic impulse, structural redistribution and NDC consistency adjustment. Use decomposition only as diagnostic accounting; document robustness tests."),
    ("2. Methods", "2.8 Interpretation and limitations", "Comparative-static production-network framework; fixed residual structure; no endogenous prices, wages, labor-market clearing, fiscal feedbacks or general-equilibrium responses."),
    ("3. Results", "3.1 Aggregate production-network effects", "Present the transition differential over time for gross output and value added relative to the preferred baseline."),
    ("3. Results", "3.2 Economic impulse vs. NDC consistency adjustment", "Compare the unconstrained investment + rewiring solution with the final NDC-constrained solution; explain where the consistency adjustment arises."),
    ("3. Results", "3.3 Sectoral winners and losers", "Identify the sectors and supply chains gaining from investment/rewiring and those contracting under the transition pathway."),
    ("3. Results", "3.4 Domestic supply-chain capture", "Show how localization/import leakage changes the domestic value-added effect and connect this result to industrial policy."),
    ("3. Results", "3.5 Labor requirements and distribution", "Report labor results only after coefficient/productivity sensitivities are finalized; keep interpretation as labor requirements under base-year productivity."),
    ("4. Discussion", "4.1–4.5 Mechanisms and implications", "Explain why low-carbon investment can offset fossil contraction; discuss domestic supply chains, NDC consistency, timing of the investment wave, trade/export assumptions and transferability."),
    ("4. Discussion", "4.6 Limitations", "Discuss IO rigidity, exogenous trajectories, price neutrality, constrained-sector residuals, labor-productivity assumptions and aggregation."),
    ("5. Conclusion", "Main conclusions", "Return to the production-network framing and identify the implications for NDC implementation and industrial policy without overstating causal or predictive claims.")
]
for secname, sub, purpose in rows:
    row = t.add_row()
    cant_split(row)
    cells = row.cells
    for j, value in enumerate([secname, sub, purpose]):
        margins(cells[j], top=70, bottom=70)
        if len(t.rows) % 2 == 0:
            shade(cells[j], "F8FAFC")
        p = cells[j].paragraphs[0]
        p.paragraph_format.line_spacing = 1.0
        font(p.add_run(value), size=8.6, bold=(j == 0), color=NAVY if j == 0 else DARK)

heading("4. Results narrative to be built around mechanisms")
para("The Results section should answer “why” the aggregate effect emerges rather than simply report endpoint totals. The preferred sequence is:", after=3)
for item in [
    "Start with the aggregate transition differential over time (output and value added).",
    "Separate the investment-driven production-network impulse from the structural redistribution caused by changes in intermediate energy demand.",
    "Show the unconstrained production-network solution and then the NDC-constrained solution, labeling the difference as the NDC consistency adjustment.",
    "Move from aggregate results to sectoral propagation: which domestic supplier chains expand, which fossil-related chains contract, and where import leakage reduces domestic capture.",
    "Use labor results as a secondary distributional layer after the employment-coefficient sensitivity is resolved."
]:
    number(item)

callout(
    "Interpretive rule for the paper",
    "The physical consistency layer must not be presented as an independent “Engine 3 shock.” It is the residual "
    "adjustment required to keep selected economic-sector outputs consistent with externally specified NDC physical "
    "trajectories after investment and technical-coefficient rewiring have propagated through the IO system.",
    fill=ORANGE
)

heading("5. Planned figures and tables")
t = doc.add_table(rows=1, cols=3)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for j, h in enumerate(["Item", "Main content", "Role in argument"]):
    c = t.cell(0, j)
    shade(c, NAVY)
    margins(c)
    font(c.paragraphs[0].add_run(h), size=9.5, bold=True, color=WHITE)

fig_rows = [
    ("Figure 1 – Model architecture", "MMA/BLUES → investment channel + energy-input rewiring → IO network, with a constraint loop from physical trajectories to selected outputs.", "Makes the two-channel + consistency-layer architecture immediately clear."),
    ("Figure 2 – How the net differential emerges", "Time profile of investment contribution, rewiring/structural redistribution, NDC consistency adjustment, and net value-added differential; optionally unconstrained vs constrained.", "Explains mechanism and timing rather than only the 2050 endpoint."),
    ("Figure 3 – Production-network winners and losers", "2050 sectoral value-added/output differences, emphasizing infrastructure, manufacturing, energy and fossil-related activities.", "Shows how aggregate effects are distributed across supply chains."),
    ("Figure 4 – Domestic capture sensitivity", "Base/low/high domestic-content or localization cases; labor distribution can be an alternative secondary panel if robust.", "Links macro results to industrial-policy implications."),
    ("Table 1 – Data and model mapping", "Main datasets, years, units, sectoral concordance and model role.", "Provides compact reproducibility overview."),
    ("Table 2 – Core assumptions / sensitivities", "Baseline, domestic content, Engine 2 mapping/valuation, physical-constraint boundaries, labor coefficients.", "Makes the interpretation and robustness framework transparent.")
]
for item, content, role in fig_rows:
    row = t.add_row()
    cant_split(row)
    for j, value in enumerate([item, content, role]):
        margins(row.cells[j], top=80, bottom=80)
        p = row.cells[j].paragraphs[0]
        p.paragraph_format.line_spacing = 1.0
        font(p.add_run(value), size=9.0, bold=(j == 0), color=NAVY if j == 0 else DARK)

heading("6. Supplementary Information (SI)")
para("ERL main text should remain compact. Detailed implementation material should move to the SI, including:", after=3)
for item in [
    "73-sector concordance and MMA/BLUES-to-IO mappings.",
    "Technology-to-sector CAPEX allocation and domestic-content coefficients.",
    "Detailed Engine 2 technical-coefficient transformations and constant-base-price interpretation.",
    "S51/S48 road-transport mapping evidence and other sector-boundary diagnostics.",
    "Cities/buildings, other-industry, aviation/maritime and remaining bespoke fuel-substitution sensitivities.",
    "Physical constraint data, constrained-sector boundary assumptions and balancing final-demand residuals.",
    "Baseline, import-leakage/localization and employment-coefficient sensitivities.",
    "Full sectoral output, value-added and labor-requirement tables, plus reproducibility notes."
]:
    bullet(item)

heading("7. Claims and terminology to preserve")
t = doc.add_table(rows=1, cols=2)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for j, h in enumerate(["Use / preferred wording", "Avoid / qualify"]):
    c = t.cell(0, j)
    shade(c, NAVY)
    margins(c)
    font(c.paragraphs[0].add_run(h), size=9.5, bold=True, color=WHITE)

use = [
    "Scenario-based, comparative-static production-network analysis.",
    "Transition/scenario differential relative to the stated counterfactual.",
    "Two economic transmission channels plus a physical consistency layer.",
    "Open Type-I Leontief system: direct and indirect supply-chain effects.",
    "NDC consistency adjustment / balancing final-demand residual.",
    "Labor requirements or job-equivalents under 2018 productivity."
]
avoid = [
    "Causal policy impact or “decarbonization causes GDP to rise by X”.",
    "Engine 3 as a third independent shock.",
    "Household-induced effects in the current scenario implementation.",
    "Employment results as forecasts of future jobs.",
    "Physical decline “destroys” a specific amount of GDP/VA.",
    "Changing physical mappings only to improve the sign of results."
]
for i in range(len(use)):
    row = t.add_row()
    cant_split(row)
    for j, values in enumerate([use, avoid]):
        c = row.cells[j]
        margins(c, top=70, bottom=70)
        shade(c, GREEN if j == 0 else "FFF2F2")
        font(c.paragraphs[0].add_run(values[i]), size=9.0)

heading("8. Current development status (September 2026)")
para(
    "The manuscript architecture and core counterfactual have been redefined and are already reflected in the paper "
    "branch. The main remaining work is to freeze the remaining model specifications before inserting final numerical "
    "results into the manuscript.",
    after=3
)
for item in [
    "Baseline: preferred all-sector GDP-growth, fixed-structure counterfactual established.",
    "Physical pathway: reframed from an independent Engine 3 shock to a hard sectoral consistency layer.",
    "Engine 2: constant-base-price interpretation documented; direct road-transport rewiring restricted to S48, with S51 treated through network linkages.",
    "Engine 2 robustness: other-industry/feedstock, cities/buildings and aviation/maritime valuation audits implemented and passing in the paper workflow.",
    "Next methodological block: refining (S19), biodiesel (S20) and gas/biomethane (S43), followed by a final audit of physical-sector boundaries.",
    "Labor: headline employment results remain secondary until the raw versus calibrated coefficient and productivity sensitivities are resolved."
]:
    bullet(item)

callout(
    "Working paper message",
    "Brazil’s NDC implementation should be analyzed as a reallocation of investment and intermediate inputs across "
    "interconnected domestic supply chains, while key sectors remain subject to physical transition constraints. "
    "The paper’s contribution is to make that reconciliation explicit and measurable."
)

props = doc.core_properties
props.title = "Overall Structure of the ERL Paper"
props.subject = "Brazil NDC production-network paper structure"
props.author = "Cenergia Lab / COPPE-UFRJ"
props.keywords = "ERL, NDC, input-output, production networks, Brazil, energy transition"

doc.save(OUT)
print(OUT)
